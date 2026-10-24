# -*- coding: utf-8 -*-
"""
Testes que rodam em qualquer sistema (sem Office e sem internet).
Executar na raiz do projeto:  python -m unittest discover -s tests
"""

import json
import os
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import traduzir_gui as tg  # noqa: E402


class MotorFake:
    """Imita GoogleTranslator.translate: traduz linha a linha prefixando 'T:'."""

    def __init__(self, juntar_linhas=False, manter=(), falhar=False):
        self.chamadas = []
        self.juntar_linhas = juntar_linhas
        self.manter = set(manter)
        self.falhar = falhar

    def translate(self, texto):
        self.chamadas.append(texto)
        if self.falhar:
            raise ConnectionError("sem rede")
        linhas = [l if l in self.manter else f"T:{l}" for l in texto.split("\n")]
        if self.juntar_linhas and len(linhas) > 1:
            return " ".join(linhas)
        return "\n".join(linhas)


def novo_tradutor(cache, motor, **kw):
    return tg.Tradutor("Português (BR)", cache, motor=motor, pausa=0, **kw)


class TestTexto(unittest.TestCase):
    def test_texto_valido(self):
        for t in ("Verificare il quadro", "SI", "NO", "N/A", "Luogo:\r\x07"):
            self.assertTrue(tg.texto_valido(t), t)
        for t in ("", " ", "A", "123", "12,5 %", "M55.780013", "A1", "PLC-01",
                  "https://exemplo.com/x", "fulano@empresa.com", None, 3.5, True):
            self.assertFalse(tg.texto_valido(t), t)

    def test_limpar_e_com_espacos(self):
        self.assertEqual(tg.limpar("  Luogo:\r\x07"), "Luogo:")
        self.assertEqual(tg.com_espacos("  abc \n", "xyz"), "  xyz \n")

    def test_nome_aba_seguro(self):
        self.assertEqual(tg.nome_aba_seguro("Está ok? [sim]", set()), "Está ok sim")
        self.assertEqual(tg.nome_aba_seguro("Lista", {"lista"}), "Lista (2)")
        self.assertEqual(len(tg.nome_aba_seguro("x" * 50, {"x" * 31})), 31)
        self.assertIsNone(tg.nome_aba_seguro("???", set()))

    def test_texto_para_celula(self):
        self.assertEqual(tg._texto_para_celula("- item"), "'- item")  # não vira fórmula
        self.assertEqual(tg._texto_para_celula("=texto"), "'=texto")
        self.assertEqual(tg._texto_para_celula("1/2"), "'1/2")  # não vira data
        self.assertEqual(tg._texto_para_celula("Texto normal"), "Texto normal")

    def test_versao(self):
        self.assertGreater(tg.versao_tupla("1.10"), tg.versao_tupla("1.9"))
        self.assertEqual(tg.versao_tupla("v1.2"), (1, 2))


class TestCache(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.arq = Path(self.dir.name) / "cache.json"

    def tearDown(self):
        self.dir.cleanup()

    def test_separado_por_idioma(self):
        c = tg.CacheManager(self.arq)
        c.definir("pt", "Hello", "Olá")
        c.definir("es", "Hello", "Hola")
        c.salvar()
        c2 = tg.CacheManager(self.arq)
        self.assertEqual(c2.obter("pt", "Hello"), "Olá")
        self.assertEqual(c2.obter("es", "Hello"), "Hola")
        self.assertIsNone(c2.obter("en", "Hello"))
        self.assertFalse(Path(str(self.arq) + ".tmp").exists())

    def test_migra_formato_antigo(self):
        self.arq.write_text(json.dumps({"Luogo:\r\x07": "Lugar:", "Vazio": None}), "utf-8")
        c = tg.CacheManager(self.arq)
        self.assertEqual(c.obter("pt", "Luogo:"), "Lugar:")
        self.assertIsNone(c.obter("pt", "Vazio"))
        c.salvar()
        self.assertEqual(json.loads(self.arq.read_text("utf-8"))["__versao__"], 2)

    def test_arquivo_corrompido_preservado(self):
        self.arq.write_text('{"quebrado": ', "utf-8")
        c = tg.CacheManager(self.arq)
        self.assertEqual(c.total(), 0)
        self.assertTrue((Path(self.dir.name) / "cache.corrompido.json").exists())


class TestTradutor(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.cache = tg.CacheManager(Path(self.dir.name) / "cache.json")

    def tearDown(self):
        self.dir.cleanup()

    def test_agrupa_requisicoes_e_usa_cache(self):
        motor = MotorFake()
        tr = novo_tradutor(self.cache, motor)
        textos = [f"Texto numero {i}" for i in range(50)]
        res = tr.traduzir_batch(textos)
        self.assertEqual(res["Texto numero 7"], "T:Texto numero 7")
        self.assertEqual(len(res), 50)
        self.assertEqual(len(motor.chamadas), 1)
        tr.traduzir_batch(textos)  # tudo no cache agora
        self.assertEqual(len(motor.chamadas), 1)

    def test_sem_agrupar(self):
        motor = MotorFake()
        novo_tradutor(self.cache, motor, agrupar=False).traduzir_batch(["Um texto", "Outro texto"])
        self.assertEqual(len(motor.chamadas), 2)

    def test_resposta_desalinhada_traduz_um_a_um(self):
        motor = MotorFake(juntar_linhas=True)
        res = novo_tradutor(self.cache, motor).traduzir_batch(["Primeiro", "Segundo", "Terceiro"])
        self.assertEqual(res, {"Primeiro": "T:Primeiro", "Segundo": "T:Segundo", "Terceiro": "T:Terceiro"})
        self.assertEqual(len(motor.chamadas), 4)

    def test_linha_inalterada_e_retentada_sozinha(self):
        motor = MotorFake(manter={"English line"})
        novo_tradutor(self.cache, motor).traduzir_batch(["Riga italiana", "English line"])
        self.assertEqual(motor.chamadas[-1], "English line")

    def test_texto_multilinha_vai_sozinho(self):
        motor = MotorFake()
        res = novo_tradutor(self.cache, motor).traduzir_batch(["Punto\nCheck-List", "Outro"])
        self.assertEqual(res["Punto\nCheck-List"], "T:Punto\nT:Check-List")
        self.assertIn("Punto\nCheck-List", motor.chamadas)

    def test_texto_longo_dividido(self):
        motor = MotorFake()
        tr = novo_tradutor(self.cache, motor)
        longo = " ".join(f"Frase numero {i}." for i in range(400))
        res = tr.traduzir_batch([longo])
        self.assertGreater(len(motor.chamadas), 1)
        self.assertTrue(all(len(c.encode()) <= tr.MAX_BYTES for c in motor.chamadas))
        self.assertEqual(res[longo].count("T:"), len(motor.chamadas))
        self.assertEqual(res[longo].replace("T:", ""), longo)

    def test_parar(self):
        stop = threading.Event()
        stop.set()
        with self.assertRaises(tg.Interrompido):
            novo_tradutor(self.cache, MotorFake(), stop_event=stop).traduzir_batch(["Algum texto"])

    def test_servico_indisponivel(self):
        tr = novo_tradutor(self.cache, MotorFake(falhar=True), agrupar=False)
        tr.ESPERA_INICIAL = 0
        with self.assertRaises(tg.ServicoIndisponivel):
            tr.traduzir_batch(["Texto um", "Texto dois", "Texto tres", "Texto quatro"])


class TestArquivos(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.raiz = self.dir.name

    def tearDown(self):
        self.dir.cleanup()

    def _criar(self, *nomes):
        for n in nomes:
            caminho = os.path.join(self.raiz, n)
            os.makedirs(os.path.dirname(caminho), exist_ok=True)
            Path(caminho).write_bytes(b"x")

    def test_escanear_ignora_saidas_e_bloqueios(self):
        self._criar("a.xls", "a_PT_BR.xlsx", "b_traduzido.xlsx", "~$c.xlsx",
                    "d.docx", "d_EN.docx", "manual_EN.docx", "notas.txt", "sub/e.doc")
        nomes = sorted(n for _, _, n, _ in tg.escanear_pasta(self.raiz, sufixo_destino="PT_BR"))
        # manual_EN.docx não tem original "manual" ao lado: é um original em inglês
        self.assertEqual(nomes, ["a.xls", "d.docx", "e.doc", "manual_EN.docx"])
        so_excel = tg.escanear_pasta(self.raiz, incluir_word=False)
        self.assertEqual([n for _, _, n, _ in so_excel], ["a.xls"])

    def test_planejar_destinos(self):
        arquivos = [
            (1, "", self.raiz, "lista.xls", "excel"),
            (2, "", self.raiz, "lista.xlsx", "excel"),
            (3, "", self.raiz, "doc.doc", "word"),
            (4, "", self.raiz, "nome.docx", "word"),
        ]
        destinos = [os.path.basename(d) for _, _, d, _ in
                    tg.planejar_destinos(arquivos, "PT_BR", {"nome": 'Nome: "traduzido"?'})]
        self.assertEqual(destinos, ["lista_PT_BR.xlsx", "lista (2)_PT_BR.xlsx",
                                    "doc_PT_BR.doc", "Nome traduzido_PT_BR.docx"])


class TestExcelOpenpyxl(unittest.TestCase):
    def test_traduz_preservando_formulas_e_tipos(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as d:
            orig, dest = os.path.join(d, "o.xlsx"), os.path.join(d, "o_PT_BR.xlsx")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Controllo"
            ws["A1"] = "Verificare il quadro"
            ws["A2"] = "=CONCATENATE(A1,\" ok\")"
            ws["A3"] = True
            ws["A4"] = 42
            ws["A5"] = "- elenco"
            ws2 = wb.create_sheet("Dati")
            ws2["A1"] = "=Controllo!A4*2"
            wb.save(orig)

            cache = tg.CacheManager(os.path.join(d, "cache.json"))
            tr = novo_tradutor(cache, MotorFake())
            ta = tg.TradutorArquivos(tr, lambda *a: None, office=None, stop_event=threading.Event())
            with mock.patch.object(tg, "TEM_COM", False):
                self.assertTrue(ta.traduzir_excel(orig, dest))

            wb = openpyxl.load_workbook(dest)
            # "Controllo" é referenciada por fórmula: não pode ser renomeada pelo openpyxl.
            # "T:Dati" perde o ":" (caractere proibido em nomes de aba).
            self.assertEqual(wb.sheetnames, ["Controllo", "TDati"])
            ws = wb["Controllo"]
            self.assertEqual(ws["A1"].value, "T:Verificare il quadro")
            self.assertEqual(ws["A2"].value, "=CONCATENATE(A1,\" ok\")")
            self.assertIs(ws["A3"].value, True)
            self.assertEqual(ws["A4"].value, 42)
            self.assertEqual(ws["A5"].value, "T:- elenco")
            self.assertEqual(wb["TDati"]["A1"].value, "=Controllo!A4*2")

    def test_xls_sem_office_falha_com_mensagem(self):
        logs = []
        tr = novo_tradutor(tg.CacheManager(os.devnull + "_inexistente"), MotorFake())
        ta = tg.TradutorArquivos(tr, lambda m, t="info": logs.append(m), None, threading.Event())
        with mock.patch.object(tg, "TEM_COM", False):
            self.assertFalse(ta.traduzir_excel("x.xls", "x_PT_BR.xlsx"))
        self.assertTrue(any(".xls" in m for m in logs))


# ---------- Objetos COM simulados (Excel/Word) ----------

class _Contagem:
    def __init__(self, n=0):
        self.Count = n


class CelulaFake:
    def __init__(self, planilha, i, j):
        self.planilha, self.i, self.j = planilha, i, j

    @property
    def Value(self):
        return self.planilha.valores[self.i - 1][self.j - 1]

    @Value.setter
    def Value(self, v):
        self.planilha.valores[self.i - 1][self.j - 1] = v


class UsedRangeFake:
    def __init__(self, planilha):
        self.planilha = planilha

    @property
    def Value(self):
        return tuple(tuple(l) for l in self.planilha.valores)

    @property
    def Formula(self):
        return tuple(
            tuple(f if f is not None else (v if v is not None else "") for f, v in zip(lf, lv))
            for lf, lv in zip(self.planilha.formulas, self.planilha.valores)
        )

    def Cells(self, i, j):
        return CelulaFake(self.planilha, i, j)


class PlanilhaFake:
    def __init__(self, nome, valores, formulas=None):
        self.Name = nome
        self.valores = [list(l) for l in valores]
        self.formulas = formulas or [[None] * len(l) for l in valores]
        self.UsedRange = UsedRangeFake(self)
        self.Shapes = []


class PastaTrabalhoFake:
    def __init__(self, planilhas):
        self.Worksheets = planilhas
        self.salvo_em = None

    def SaveAs(self, Filename, FileFormat):
        self.salvo_em = (Filename, FileFormat)

    def Close(self, SaveChanges=False):
        pass


class ExcelFake:
    def __init__(self, wb):
        self.Calculation = -4105
        self.calculos = []
        wb_ref = wb

        class Workbooks:
            @staticmethod
            def Open(**kw):
                return wb_ref
        self.Workbooks = Workbooks


class RangeWordFake:
    duplicados = []

    def __init__(self, texto, objetos=0):
        self.Text = texto
        self.Fields = _Contagem(objetos)
        self.InlineShapes = self.ShapeRange = self.Footnotes = _Contagem()
        self.Endnotes = self.ContentControls = _Contagem()

    @property
    def Duplicate(self):
        dup = RangeWordFake(self.Text, self.Fields.Count)
        RangeWordFake.duplicados.append(dup)
        return dup

    def MoveEnd(self, unidade, n):
        marca = 2 if self.Text.endswith("\r\x07") else 1
        self.Text = self.Text[:-marca]


class ParagrafoFake:
    def __init__(self, texto, objetos=0):
        self.Range = RangeWordFake(texto, objetos)


class StoryFake:
    def __init__(self, paragrafos):
        self.Paragraphs = paragrafos
        self.NextStoryRange = None


class DocumentoFake:
    def __init__(self, paragrafos):
        self.StoryRanges = [StoryFake(paragrafos)]
        self.TrackRevisions = True
        self.salvo_em = None

    def SaveAs(self, caminho, formato):
        self.salvo_em = (caminho, formato)

    def Close(self, salvar):
        pass


class OfficeFake:
    def __init__(self, app):
        self.app = app

    def obter(self, progid):
        return self.app


class TestComSimulado(unittest.TestCase):
    def _arquivos(self, app):
        d = tempfile.mkdtemp()
        self.addCleanup(lambda: __import__("shutil").rmtree(d, ignore_errors=True))
        tr = novo_tradutor(tg.CacheManager(os.path.join(d, "c.json")), MotorFake())
        logs = []
        ta = tg.TradutorArquivos(tr, lambda m, t="info": logs.append(m), OfficeFake(app), threading.Event())
        return ta, d, logs

    def test_excel(self):
        ws = PlanilhaFake("Foglio?", [
            ["Verificare il quadro", 3.5, "Totale"],
            [None, "- elenco puntato", "  con spazi "],
        ], formulas=[[None, None, '=A1&" ok"'], [None, None, None]])
        outra = PlanilhaFake("Fogliox", [["Nota finale"]])
        wb = PastaTrabalhoFake([ws, outra])
        excel = ExcelFake(wb)
        ta, d, _ = self._arquivos(excel)
        with mock.patch.object(tg, "TEM_COM", True):
            self.assertTrue(ta.traduzir_excel("a.xls", os.path.join(d, "a_PT_BR.xlsx")))
        self.assertEqual(ws.valores[0], ["T:Verificare il quadro", 3.5, "Totale"])  # fórmula intacta
        self.assertEqual(ws.valores[1][1], "T:- elenco puntato")
        self.assertEqual(ws.valores[1][2], "  T:con spazi ")
        self.assertEqual(ws.Name, "TFoglio")  # ':' e '?' removidos
        self.assertEqual(excel.Calculation, -4105)  # modo de cálculo restaurado
        self.assertEqual(wb.salvo_em[1], 51)

    def test_word(self):
        paragrafos = [
            ParagrafoFake("Titolo del documento\r"),
            ParagrafoFake("Luogo:\r\x07"),
            ParagrafoFake("Prima riga\x0bSeconda riga\r"),
            ParagrafoFake("Pagina con campo\r", objetos=1),
            ParagrafoFake("123\r"),
        ]
        doc = DocumentoFake(paragrafos)
        RangeWordFake.duplicados = []

        class Documents:
            @staticmethod
            def Open(**kw):
                return doc
        word = type("WordFake", (), {"Documents": Documents})()
        ta, d, logs = self._arquivos(word)
        with mock.patch.object(tg, "TEM_COM", True):
            self.assertTrue(ta.traduzir_word("a.docx", os.path.join(d, "a_PT_BR.docx")))
        # O texto novo é gravado no Range duplicado (sem a marca de parágrafo)
        textos = [r.Text for r in RangeWordFake.duplicados]
        self.assertIn("T:Titolo del documento", textos)
        self.assertIn("T:Luogo:", textos)
        self.assertIn("T:Prima riga\x0bT:Seconda riga", textos)
        self.assertIn("Pagina con campo", textos)  # mantido no original
        self.assertTrue(any("imagens/campos" in m for m in logs))
        self.assertFalse(doc.TrackRevisions)
        self.assertEqual(doc.salvo_em[1], 16)


if __name__ == "__main__":
    unittest.main()
